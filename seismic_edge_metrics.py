"""
Seismic edge/topology metric comparison for SLoG, LoG, and Canny.

The script accepts either one seismic profile image or a folder of images. For
all images it applies the same preprocessing pipeline, runs the three detectors,
computes topology/continuity/spatial-correspondence metrics, and writes:

1. seismic_edge_metrics.xlsx
   - per_image: one row per image and algorithm
   - algorithm_summary: mean/std/median per algorithm
   - metric_definitions: metric definitions and interpretation
   - parameters: all runtime parameters
2. seismic_edge_metrics.csv
3. algorithm_summary.csv
4. optional intermediate/edge maps under output/maps/

Example:
    python seismic_edge_metrics.py \
        --input /path/to/images \
        --output /path/to/results \
        --save-maps

Optional fault masks:
    python seismic_edge_metrics_compare.py \
        --input /path/to/images \
        --fault-mask-dir /path/to/fault_masks \
        --output /path/to/results

Mask matching is by filename stem. A nonzero mask pixel is treated as fault.

Default source-compatible parameters:
    SGS: sigma=1.5, anisotropy=2.2, iterations=1, base_kappa=10.0,
         gamma=0.06, nl_h=0.8
    AGC: window=31, eps=1e-8, p1=1
    SLoG: sigma=0.4, anisotropy=2.1, n_angles=8, sharpness=14,
          alpha=0.92, aniso_clip=4, gamma_scale=1, dist_mode=relative
    LoG: sigma=1.1
    Canny: sigma=1.7, low=0.1, high=0.9

Core dependencies:
    numpy pandas scipy pillow scikit-image openpyxl tqdm
Optional:
    torch (used to accelerate the structure-guided diffusion; a NumPy/SciPy
    fallback is included when torch is unavailable)
"""

from __future__ import annotations

import argparse
import json
import logging
import math
import os
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any, Dict, Iterable, List, Mapping, Optional, Sequence, Tuple

import numpy as np
import pandas as pd
from PIL import Image
from scipy import ndimage
from scipy.ndimage import gaussian_filter, uniform_filter
from skimage import feature, filters, measure, morphology, restoration, img_as_ubyte
from tqdm import tqdm

try:
    import torch
    import torch.nn.functional as F
except Exception:  # pragma: no cover - optional dependency
    torch = None
    F = None


EPS = 1e-12
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}

DEFAULT_SGS_PARAMS: Dict[str, Any] = {
    "sigma": 1.5,
    "anisotropy": 2.2,
    "iterations": 1,
    "base_kappa": 10.0,
    "gamma": 0.06,
    "device": "auto",
    "prefer_fp16": False,
    "nl_h": 0.8,
    "option": 1,
    "lambda_perp_base": 1e-3,
    "lambda_par_base": 1.0,
    "coherence_influence": 1.0,
}

DEFAULT_AGC_PARAMS: Dict[str, Any] = {
    "window": 31,
    "eps": 1e-8,
    "p1": 1,
}

DEFAULT_DET_PARAMS: Dict[str, Any] = {
    "tex_layout": "fault,agc,edge",
    "glog_sigma": 0.4,
    "glog_anisotropy": 2.1,
    "glog_nangles": 8,
    "glog_sharpness": 14.0,
    "glog_alpha": 0.92,
    "glog_aniso_clip": 4.0,
    "glog_gamma_scale": 1.0,
    "glog_dist_mode": "relative",
    "glog_coherence_influence": 1.0,
    "glog_size_factor": 6.0,
    "glog_use_absolute_sigma": False,
    "edge_thr": 0.35,
    "log_sigma": 1.1,
    "canny_sigma": 1.7,
    "canny_low": 0.1,
    "canny_high": 0.9,
}


@dataclass
class PreprocessConfig:
    """Shared preprocessing for all three edge detectors."""

    nl_h: float = float(DEFAULT_SGS_PARAMS["nl_h"])
    nlmeans_patch_size: int = 5
    nlmeans_patch_distance: int = 6

    sgs_sigma: float = float(DEFAULT_SGS_PARAMS["sigma"])
    sgs_anisotropy: float = float(DEFAULT_SGS_PARAMS["anisotropy"])
    sgs_iterations: int = int(DEFAULT_SGS_PARAMS["iterations"])
    sgs_base_kappa: float = float(DEFAULT_SGS_PARAMS["base_kappa"])
    sgs_gamma: float = float(DEFAULT_SGS_PARAMS["gamma"])
    sgs_option: int = int(DEFAULT_SGS_PARAMS["option"])
    lambda_perp_base: float = float(DEFAULT_SGS_PARAMS["lambda_perp_base"])
    lambda_par_base: float = float(DEFAULT_SGS_PARAMS["lambda_par_base"])
    coherence_influence: float = float(DEFAULT_SGS_PARAMS["coherence_influence"])

    agc_window: int = int(DEFAULT_AGC_PARAMS["window"])
    agc_eps: float = float(DEFAULT_AGC_PARAMS["eps"])
    agc_p1: float = float(DEFAULT_AGC_PARAMS["p1"])

    percentile_low: float = 1.0
    percentile_high: float = 99.0
    device: str = str(DEFAULT_SGS_PARAMS["device"])
    prefer_fp16: bool = bool(DEFAULT_SGS_PARAMS["prefer_fp16"])


@dataclass
class DetectorConfig:
    """Parameters for SLoG, LoG and Canny."""

    tex_layout: str = str(DEFAULT_DET_PARAMS["tex_layout"])

    glog_sigma: float = float(DEFAULT_DET_PARAMS["glog_sigma"])
    glog_anisotropy: float = float(DEFAULT_DET_PARAMS["glog_anisotropy"])
    glog_nangles: int = int(DEFAULT_DET_PARAMS["glog_nangles"])
    glog_size_factor: float = float(DEFAULT_DET_PARAMS["glog_size_factor"])
    glog_sharpness: float = float(DEFAULT_DET_PARAMS["glog_sharpness"])
    glog_alpha: float = float(DEFAULT_DET_PARAMS["glog_alpha"])
    glog_aniso_clip: float = float(DEFAULT_DET_PARAMS["glog_aniso_clip"])
    glog_gamma_scale: float = float(DEFAULT_DET_PARAMS["glog_gamma_scale"])
    glog_dist_mode: str = str(DEFAULT_DET_PARAMS["glog_dist_mode"])
    glog_coherence_influence: float = float(DEFAULT_DET_PARAMS["glog_coherence_influence"])
    glog_use_absolute_sigma: bool = bool(DEFAULT_DET_PARAMS["glog_use_absolute_sigma"])
    glog_absolute_response: bool = False
    edge_thr: float = float(DEFAULT_DET_PARAMS["edge_thr"])

    log_sigma: float = float(DEFAULT_DET_PARAMS["log_sigma"])
    log_edge_threshold: float = -1.0

    canny_sigma: float = float(DEFAULT_DET_PARAMS["canny_sigma"])
    canny_low: float = float(DEFAULT_DET_PARAMS["canny_low"])
    canny_high: float = float(DEFAULT_DET_PARAMS["canny_high"])
    canny_use_uint8_input: bool = True

    use_response_hysteresis: bool = False
    hysteresis_low_ratio: float = 0.50


@dataclass
class MetricConfig:
    """Parameters controlling metric definitions."""

    short_edge_length_px: float = 10.0
    isolated_max_length_px: float = 2.0
    continuity_lookahead_px: int = 7
    continuity_tolerance_px: int = 1
    orientation_bins: int = 18
    orientation_entropy_window: int = 21
    orientation_entropy_min_edges: int = 4
    low_coherence_percentile: float = 20.0
    low_coherence_tolerance_px: int = 2
    fault_tolerance_px: int = 3
    valid_local_window: int = 15
    valid_std_percentile: float = 10.0


# -----------------------------------------------------------------------------
# General utilities
# -----------------------------------------------------------------------------


def ensure_dir(path: Path | str) -> Path:
    p = Path(path)
    p.mkdir(parents=True, exist_ok=True)
    return p


def normalize_map(arr: np.ndarray) -> np.ndarray:
    arr = np.asarray(arr, dtype=np.float32)
    arr = np.nan_to_num(arr, nan=0.0, posinf=0.0, neginf=0.0)
    lo = float(np.min(arr))
    hi = float(np.max(arr))
    if hi - lo <= EPS:
        return np.zeros_like(arr, dtype=np.float32)
    return ((arr - lo) / (hi - lo)).astype(np.float32)


def percentile_normalize(arr: np.ndarray, low: float = 1.0, high: float = 99.0) -> np.ndarray:
    arr = np.asarray(arr, dtype=np.float32)
    arr = np.nan_to_num(arr, nan=0.0, posinf=0.0, neginf=0.0)
    lo, hi = np.percentile(arr, [low, high])
    if not np.isfinite(lo) or not np.isfinite(hi) or hi - lo <= EPS:
        return normalize_map(arr)
    return np.clip((arr - lo) / (hi - lo + EPS), 0.0, 1.0).astype(np.float32)


def load_grayscale_image(path: Path, cfg: PreprocessConfig) -> np.ndarray:
    """Load 8/16-bit grayscale or RGB image and robustly normalize to [0, 1]."""
    with Image.open(path) as im:
        arr = np.asarray(im)

    if arr.ndim == 3:
        if arr.shape[-1] >= 3:
            arr = (
                0.2989 * arr[..., 0].astype(np.float32)
                + 0.5870 * arr[..., 1].astype(np.float32)
                + 0.1140 * arr[..., 2].astype(np.float32)
            )
        else:
            arr = arr[..., 0].astype(np.float32)
    else:
        arr = arr.astype(np.float32)

    return percentile_normalize(arr, cfg.percentile_low, cfg.percentile_high)


def save_gray(path: Path, arr: np.ndarray) -> None:
    a = np.clip(np.nan_to_num(arr, nan=0.0), 0.0, 1.0)
    Image.fromarray(np.round(a * 255).astype(np.uint8), mode="L").save(path)


def save_binary(path: Path, arr: np.ndarray) -> None:
    Image.fromarray((np.asarray(arr, dtype=bool) * 255).astype(np.uint8), mode="L").save(path)


def save_edge_overlay(path: Path, background: np.ndarray, edge: np.ndarray) -> None:
    bg = np.round(np.clip(background, 0.0, 1.0) * 255).astype(np.uint8)
    rgb = np.stack([bg, bg, bg], axis=-1)
    e = np.asarray(edge, dtype=bool)
    rgb[e, 0] = 255
    rgb[e, 1] = 0
    rgb[e, 2] = 0
    Image.fromarray(rgb, mode="RGB").save(path)


def collect_image_paths(input_path: Path, recursive: bool) -> List[Path]:
    if input_path.is_file():
        if input_path.suffix.lower() not in IMAGE_EXTENSIONS:
            raise ValueError(f"Unsupported image type: {input_path}")
        return [input_path]
    if not input_path.is_dir():
        raise FileNotFoundError(f"Input path not found: {input_path}")
    iterator = input_path.rglob("*") if recursive else input_path.glob("*")
    return sorted(p for p in iterator if p.is_file() and p.suffix.lower() in IMAGE_EXTENSIONS)


def odd_at_least_three(value: int) -> int:
    value = max(3, int(value))
    return value if value % 2 == 1 else value + 1


# -----------------------------------------------------------------------------
# Preprocessing: NL-means -> structure-guided smoothing -> AGC
# -----------------------------------------------------------------------------


def denoise_nlmeans(img: np.ndarray, cfg: PreprocessConfig) -> np.ndarray:
    try:
        try:
            sigma_est = restoration.estimate_sigma(img, channel_axis=None)
        except TypeError:  
            sigma_est = restoration.estimate_sigma(img, multichannel=False)
        sigma_val = float(np.mean(sigma_est))
    except Exception:
        high_pass = img.astype(np.float32) - gaussian_filter(img.astype(np.float32), sigma=1.0)
        median = float(np.median(high_pass))
        sigma_val = float(np.median(np.abs(high_pass - median)) / 0.67448975)
    h = max(float(cfg.nl_h) * sigma_val, 1e-6)
    kwargs = dict(
        h=h,
        patch_size=int(cfg.nlmeans_patch_size),
        patch_distance=int(cfg.nlmeans_patch_distance),
        fast_mode=True,
        preserve_range=True,
    )
    try:
        den = restoration.denoise_nl_means(img, channel_axis=None, **kwargs)
    except TypeError:
        den = restoration.denoise_nl_means(img, multichannel=False, **kwargs)
    return np.clip(den, 0.0, 1.0).astype(np.float32)


def build_structure_tensor_maps(
    img: np.ndarray,
    sigma: float,
    lambda_parallel: float,
    lambda_perp: float,
    coherence_influence: float,
) -> Tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    """Return coherence, principal orientation and diffusion tensor components."""
    gx = ndimage.sobel(img, axis=1, mode="reflect").astype(np.float32)
    gy = ndimage.sobel(img, axis=0, mode="reflect").astype(np.float32)
    jxx = gaussian_filter(gx * gx, sigma=sigma, mode="reflect")
    jxy = gaussian_filter(gx * gy, sigma=sigma, mode="reflect")
    jyy = gaussian_filter(gy * gy, sigma=sigma, mode="reflect")

    delta = np.sqrt((jxx - jyy) ** 2 + 4.0 * jxy * jxy)
    eig1 = 0.5 * (jxx + jyy + delta)
    eig2 = 0.5 * (jxx + jyy - delta)
    coherence = np.clip((eig1 - eig2) / (eig1 + eig2 + EPS), 0.0, 1.0)
    orientation = 0.5 * np.arctan2(2.0 * jxy, jxx - jyy + EPS)

    vx = np.cos(orientation)
    vy = np.sin(orientation)
    lam_per_map = float(lambda_perp) * (1.0 - float(coherence_influence) * coherence)
    lam_per_map = np.clip(lam_per_map, 1e-8, None)
    lam_par = float(lambda_parallel)
    dxx = lam_par * vx * vx + lam_per_map * vy * vy
    dyy = lam_par * vy * vy + lam_per_map * vx * vx
    dxy = (lam_par - lam_per_map) * vx * vy

    return (
        coherence.astype(np.float32),
        orientation.astype(np.float32),
        dxx.astype(np.float32),
        dxy.astype(np.float32),
        dyy.astype(np.float32),
    )


def _resolve_torch_device(device_name: str):
    if torch is None:
        return None
    name = str(device_name).lower()
    if name == "auto":
        return torch.device("cuda:0" if torch.cuda.is_available() else "cpu")
    if name.startswith("cuda") and not torch.cuda.is_available():
        logging.warning("CUDA requested but unavailable; using CPU.")
        return torch.device("cpu")
    return torch.device(device_name)


def tensor_diffusion_torch(
    img: np.ndarray,
    dxx: np.ndarray,
    dxy: np.ndarray,
    dyy: np.ndarray,
    iterations: int,
    gamma: float,
    device_name: str,
    prefer_fp16: bool = False,
) -> np.ndarray:
    if torch is None or F is None:
        raise RuntimeError("PyTorch is unavailable")
    device = _resolve_torch_device(device_name)
    use_fp16 = bool(prefer_fp16) and device.type == "cuda"
    dtype = torch.float16 if use_fp16 else torch.float32
    image = torch.from_numpy(img.astype(np.float32))[None, None].to(device=device, dtype=dtype)
    tdxx = torch.from_numpy(dxx)[None, None].to(device=device, dtype=dtype)
    tdxy = torch.from_numpy(dxy)[None, None].to(device=device, dtype=dtype)
    tdyy = torch.from_numpy(dyy)[None, None].to(device=device, dtype=dtype)
    kx = torch.tensor([[[[-0.5, 0.0, 0.5]]]], device=device, dtype=dtype)
    ky = torch.tensor([[[[-0.5], [0.0], [0.5]]]], device=device, dtype=dtype)

    out = image
    for _ in range(max(0, int(iterations))):
        gx = F.conv2d(F.pad(out, (1, 1, 0, 0), mode="replicate"), kx)
        gy = F.conv2d(F.pad(out, (0, 0, 1, 1), mode="replicate"), ky)
        jx = tdxx * gx + tdxy * gy
        jy = tdxy * gx + tdyy * gy
        djx = F.conv2d(F.pad(jx, (1, 1, 0, 0), mode="replicate"), kx)
        djy = F.conv2d(F.pad(jy, (0, 0, 1, 1), mode="replicate"), ky)
        out = torch.nan_to_num(out + float(gamma) * (djx + djy), nan=0.0, posinf=1e4, neginf=-1e4)

    return out[0, 0].detach().cpu().numpy().astype(np.float32)


def tensor_diffusion_numpy(
    img: np.ndarray,
    dxx: np.ndarray,
    dxy: np.ndarray,
    dyy: np.ndarray,
    iterations: int,
    gamma: float,
) -> np.ndarray:
    out = img.astype(np.float32).copy()
    for _ in range(max(0, int(iterations))):
        gx = ndimage.convolve1d(out, np.array([-0.5, 0.0, 0.5], dtype=np.float32), axis=1, mode="nearest")
        gy = ndimage.convolve1d(out, np.array([-0.5, 0.0, 0.5], dtype=np.float32), axis=0, mode="nearest")
        jx = dxx * gx + dxy * gy
        jy = dxy * gx + dyy * gy
        djx = ndimage.convolve1d(jx, np.array([-0.5, 0.0, 0.5], dtype=np.float32), axis=1, mode="nearest")
        djy = ndimage.convolve1d(jy, np.array([-0.5, 0.0, 0.5], dtype=np.float32), axis=0, mode="nearest")
        out = np.nan_to_num(out + float(gamma) * (djx + djy), nan=0.0, posinf=1e4, neginf=-1e4)
    return out.astype(np.float32)


def structure_guided_smoothing(
    img: np.ndarray,
    sigma: float = 1.5,
    anisotropy: float = 2.2,
    iterations: int = 1,
    base_kappa: float = 10.0,
    gamma: float = 0.06,
    option: int = 1,
    lambda_perp_base: float = 1e-3,
    lambda_par_base: float = 1.0,
    coherence_influence: float = 1.0,
    device: str = "auto",
    prefer_fp16: bool = False,
) -> Tuple[np.ndarray, np.ndarray, np.ndarray, np.ndarray, np.ndarray, np.ndarray]:
    """SGS with the complete original parameter signature.

    The supplied source accepts ``anisotropy``, ``base_kappa`` and ``option``
    but does not use them in its current equations. They are now passed into
    this function explicitly rather than merely stored in a config object.
    """
    sigma = float(sigma)
    anisotropy = float(anisotropy)
    iterations = int(iterations)
    base_kappa = float(base_kappa)
    gamma = float(gamma)
    option = int(option)
    lambda_perp_base = float(lambda_perp_base)
    lambda_par_base = float(lambda_par_base)
    coherence_influence = float(coherence_influence)

    if sigma <= 0:
        raise ValueError(f"sgs sigma must be > 0, got {sigma}")
    if anisotropy <= 0:
        raise ValueError(f"sgs anisotropy must be > 0, got {anisotropy}")
    if iterations < 0:
        raise ValueError(f"sgs iterations must be >= 0, got {iterations}")
    if base_kappa <= 0:
        raise ValueError(f"sgs base_kappa must be > 0, got {base_kappa}")

    coherence, orientation, dxx, dxy, dyy = build_structure_tensor_maps(
        img,
        sigma=sigma,
        lambda_parallel=lambda_par_base,
        lambda_perp=lambda_perp_base,
        coherence_influence=coherence_influence,
    )
    if torch is not None:
        try:
            smooth = tensor_diffusion_torch(
                img,
                dxx,
                dxy,
                dyy,
                iterations=iterations,
                gamma=gamma,
                device_name=device,
                prefer_fp16=prefer_fp16,
            )
        except Exception as exc:
            logging.warning("Torch diffusion failed (%s); using NumPy fallback.", exc)
            smooth = tensor_diffusion_numpy(img, dxx, dxy, dyy, iterations, gamma)
    else:
        smooth = tensor_diffusion_numpy(img, dxx, dxy, dyy, iterations, gamma)

    smooth = normalize_map(smooth)
    return smooth, coherence, orientation, dxx, dxy, dyy


def agc_local_rms(img: np.ndarray, window: int, low_percentile: float, eps: float = 1e-8) -> np.ndarray:
    window = odd_at_least_three(window)
    local_mean_sq = uniform_filter(img.astype(np.float32) ** 2, size=window, mode="reflect")
    local_rms = np.sqrt(np.maximum(local_mean_sq, 0.0) + eps)
    out = img / (local_rms + eps)
    p = float(np.clip(low_percentile, 0.0, 49.9))
    lo, hi = np.percentile(out, [p, 100.0 - p])
    if hi - lo <= EPS:
        return normalize_map(out)
    return np.clip((out - lo) / (hi - lo + EPS), 0.0, 1.0).astype(np.float32)


def preprocess_image(img: np.ndarray, cfg: PreprocessConfig) -> Dict[str, np.ndarray]:
    nl = denoise_nlmeans(img, cfg)
    sgs, coherence, orientation, dxx, dxy, dyy = structure_guided_smoothing(
        nl,
        sigma=cfg.sgs_sigma,
        anisotropy=cfg.sgs_anisotropy,
        iterations=cfg.sgs_iterations,
        base_kappa=cfg.sgs_base_kappa,
        gamma=cfg.sgs_gamma,
        option=cfg.sgs_option,
        lambda_perp_base=cfg.lambda_perp_base,
        lambda_par_base=cfg.lambda_par_base,
        coherence_influence=cfg.coherence_influence,
        device=cfg.device,
        prefer_fp16=cfg.prefer_fp16,
    )
    agc = agc_local_rms(sgs, cfg.agc_window, cfg.agc_p1, eps=cfg.agc_eps)
    return {
        "gray": img.astype(np.float32),
        "nlmeans": nl,
        "sgs": sgs,
        "agc": agc,
        "coherence": coherence,
        "orientation": orientation,
        "Dxx": dxx,
        "Dxy": dxy,
        "Dyy": dyy,
    }


# -----------------------------------------------------------------------------
# SLoG, LoG and Canny
# -----------------------------------------------------------------------------


def make_anisotropic_gaussian_kernel(sx: float, sy: float, angle_rad: float, size: int) -> np.ndarray:
    half = size // 2
    y, x = np.mgrid[-half : half + 1, -half : half + 1].astype(np.float32)
    ca = math.cos(angle_rad)
    sa = math.sin(angle_rad)
    xr = ca * x + sa * y
    yr = -sa * x + ca * y
    g = np.exp(-0.5 * ((xr / (sx + EPS)) ** 2 + (yr / (sy + EPS)) ** 2))
    return (g / (np.sum(g) + EPS)).astype(np.float32)


def laplacian_of_kernel(kernel: np.ndarray) -> np.ndarray:
    lap = ndimage.laplace(kernel, mode="reflect")
    lap = lap - np.mean(lap)
    return (lap / (np.sum(np.abs(lap)) + EPS)).astype(np.float32)


def slog_response(
    img: np.ndarray,
    orientation_map: np.ndarray,
    coherence_map: np.ndarray,
    dxx_map: np.ndarray,
    dxy_map: np.ndarray,
    dyy_map: np.ndarray,
    cfg: DetectorConfig,
) -> np.ndarray:
    """
    Structure-tensor-guided smoothing-based LoG response.

    This follows the supplied source file's generalized oriented LoG: an
    anisotropic LoG filter bank is weighted by tensor orientation, coherence,
    local anisotropy and angular alignment, then fused pixelwise.
    """
    image = img.astype(np.float32)
    angles = np.linspace(0.0, np.pi, int(cfg.glog_nangles), endpoint=False)
    max_aniso = max(1.0, float(cfg.glog_anisotropy), float(cfg.glog_aniso_clip))
    kernel_size = int(max(3, math.ceil(cfg.glog_size_factor * cfg.glog_sigma * max_aniso) * 2 + 1))
    if kernel_size % 2 == 0:
        kernel_size += 1

    a = dxx_map.astype(np.float32)
    b = dxy_map.astype(np.float32)
    c = dyy_map.astype(np.float32)
    delta = np.sqrt((a - c) ** 2 + 4.0 * b * b + 1e-18)
    eig1 = np.clip(0.5 * (a + c + delta), EPS, None)
    eig2 = np.clip(0.5 * (a + c - delta), EPS, None)
    tensor_orientation = 0.5 * np.arctan2(2.0 * b, a - c + EPS)
    anisotropy_ratio = np.clip(np.sqrt(eig1 / eig2), 1.0, float(cfg.glog_aniso_clip))
    ratio_scaled = np.power(anisotropy_ratio, float(cfg.glog_gamma_scale))
    sigma_parallel = float(cfg.glog_sigma) * ratio_scaled
    sigma_perpendicular = float(cfg.glog_sigma) / (ratio_scaled + EPS)
    if cfg.glog_use_absolute_sigma:
        sigma_parallel = np.abs(sigma_parallel)
        sigma_perpendicular = np.abs(sigma_perpendicular)

    fused = np.zeros_like(image, dtype=np.float32)
    if cfg.glog_dist_mode.lower() == "relative":
        max_alignment = np.zeros_like(image, dtype=np.float32)
        for angle in angles:
            max_alignment = np.maximum(max_alignment, np.abs(np.cos(float(angle) - tensor_orientation)))
    else:
        max_alignment = np.ones_like(image, dtype=np.float32)

    for angle in angles:
        kernel = make_anisotropic_gaussian_kernel(
            sx=float(cfg.glog_sigma) * float(cfg.glog_anisotropy),
            sy=float(cfg.glog_sigma),
            angle_rad=float(angle),
            size=kernel_size,
        )
        log_kernel = laplacian_of_kernel(kernel)
        response = ndimage.convolve(image, log_kernel, mode="reflect").astype(np.float32)

        alignment = np.abs(np.cos(float(angle) - tensor_orientation))
        if cfg.glog_dist_mode.lower() == "relative":
            alignment = alignment / (max_alignment + 1e-6)
        if float(cfg.glog_sharpness) != 1.0:
            alignment = np.power(alignment, float(cfg.glog_sharpness))

        coherence_weight = np.clip(
            float(cfg.glog_coherence_influence) * coherence_map, 0.0, 1.0
        )
        base_weight = (
            (1.0 - float(cfg.glog_alpha))
            + float(cfg.glog_alpha) * coherence_weight * alignment
        )
        local_scale_ratio = sigma_parallel / (float(cfg.glog_sigma) + EPS)
        anisotropy_boost = 1.0 + (local_scale_ratio - 1.0) * alignment
        fused += response * base_weight * anisotropy_boost

    local_scale = np.sqrt(np.maximum(sigma_parallel, EPS) * np.maximum(sigma_perpendicular, EPS)) / (
        float(cfg.glog_sigma) + EPS
    )
    fused *= np.clip(local_scale, 1e-6, 1e3)
    if cfg.glog_absolute_response:
        fused = np.abs(fused)
    fused = normalize_map(fused)
    fused = percentile_normalize(fused, low=5.0, high=99.0)
    return fused.astype(np.float32)


def log_response(img: np.ndarray, sigma: float) -> np.ndarray:
    smooth = gaussian_filter(img.astype(np.float32), sigma=float(sigma), mode="reflect")
    response = np.abs(ndimage.laplace(smooth, mode="reflect"))
    return normalize_map(response)


def response_to_binary(
    response: np.ndarray,
    threshold: float,
    hysteresis_low_ratio: float,
    use_hysteresis: bool,
) -> Tuple[np.ndarray, float]:
    values = response[np.isfinite(response) & (response > 0)]
    if values.size == 0:
        return np.zeros_like(response, dtype=bool), float("nan")
    if threshold < 0:
        try:
            high = float(filters.threshold_otsu(values))
        except Exception:
            high = float(np.percentile(values, 75.0))
    else:
        high = float(threshold)
    high = float(np.clip(high, 0.0, 1.0))
    if use_hysteresis:
        low = float(np.clip(hysteresis_low_ratio * high, 0.0, high))
        binary = filters.apply_hysteresis_threshold(response, low, high)
    else:
        binary = response >= high
    return np.asarray(binary, dtype=bool), high


def run_detectors(pre: Mapping[str, np.ndarray], cfg: DetectorConfig) -> Dict[str, Dict[str, Any]]:
    agc = pre["agc"]
    sr = slog_response(
        agc,
        pre["orientation"],
        pre["coherence"],
        pre["Dxx"],
        pre["Dxy"],
        pre["Dyy"],
        cfg,
    )
    sb, st = response_to_binary(
        sr, cfg.edge_thr, cfg.hysteresis_low_ratio, cfg.use_response_hysteresis
    )

    lr = log_response(agc, cfg.log_sigma)
    lb, lt = response_to_binary(
        lr, cfg.log_edge_threshold, cfg.hysteresis_low_ratio, cfg.use_response_hysteresis
    )

    canny_input = img_as_ubyte(np.clip(agc, 0.0, 1.0)) if cfg.canny_use_uint8_input else agc.astype(np.float32)
    cb = feature.canny(
        canny_input,
        sigma=float(cfg.canny_sigma),
        low_threshold=float(cfg.canny_low),
        high_threshold=float(cfg.canny_high),
    )
    cr = cb.astype(np.float32)

    return {
        "SLoG": {"response": sr, "binary": sb, "threshold": st},
        "LoG": {"response": lr, "binary": lb, "threshold": lt},
        "Canny": {"response": cr, "binary": cb.astype(bool), "threshold": float(cfg.canny_high)},
    }


# -----------------------------------------------------------------------------
# Metric computation
# -----------------------------------------------------------------------------


def make_valid_signal_mask(img: np.ndarray, cfg: MetricConfig) -> np.ndarray:
    win = odd_at_least_three(cfg.valid_local_window)
    mean = uniform_filter(img, size=win, mode="reflect")
    mean2 = uniform_filter(img * img, size=win, mode="reflect")
    local_std = np.sqrt(np.maximum(mean2 - mean * mean, 0.0))
    positive = local_std[local_std > 1e-7]
    if positive.size == 0:
        return np.ones_like(img, dtype=bool)
    threshold = max(float(np.percentile(positive, cfg.valid_std_percentile)), 1e-5)
    mask = local_std >= threshold
    mask = morphology.closing(mask, morphology.disk(2))
    try:
        mask = morphology.remove_small_holes(mask, max_size=63)
    except TypeError: 
        mask = morphology.remove_small_holes(mask, area_threshold=64)
    return np.asarray(mask, dtype=bool)


def graph_length_of_component(component: np.ndarray) -> float:
    component = np.asarray(component, dtype=bool)
    horizontal = np.count_nonzero(component[:, :-1] & component[:, 1:])
    vertical = np.count_nonzero(component[:-1, :] & component[1:, :])
    diagonal_1 = np.count_nonzero(component[:-1, :-1] & component[1:, 1:])
    diagonal_2 = np.count_nonzero(component[1:, :-1] & component[:-1, 1:])
    length = float(horizontal + vertical + math.sqrt(2.0) * (diagonal_1 + diagonal_2))
    return max(1.0, length)


def component_length_data(skeleton: np.ndarray) -> Tuple[np.ndarray, np.ndarray, np.ndarray]:
    labels = measure.label(skeleton, connectivity=2)
    lengths: List[float] = []
    pixel_counts: List[int] = []
    for label_id in range(1, int(labels.max()) + 1):
        component = labels == label_id
        lengths.append(graph_length_of_component(component))
        pixel_counts.append(int(np.count_nonzero(component)))
    return labels, np.asarray(lengths, dtype=np.float64), np.asarray(pixel_counts, dtype=np.int64)


def skeleton_degree(skeleton: np.ndarray) -> np.ndarray:
    kernel = np.ones((3, 3), dtype=np.uint8)
    kernel[1, 1] = 0
    return ndimage.convolve(skeleton.astype(np.uint8), kernel, mode="constant", cval=0)


def count_junction_clusters(skeleton: np.ndarray) -> Tuple[int, int, int]:
    degree = skeleton_degree(skeleton)
    candidates = skeleton & (degree >= 3)
    labels = measure.label(candidates, connectivity=2)
    branch_count = 0
    crossing_count = 0
    total = 0
    for label_id in range(1, int(labels.max()) + 1):
        cluster = labels == label_id
        total += 1
        if int(np.max(degree[cluster])) >= 4:
            crossing_count += 1
        else:
            branch_count += 1
    return branch_count, crossing_count, total


def local_tangent_orientation(pre_orientation: np.ndarray) -> np.ndarray:
    return np.mod(pre_orientation + np.pi / 2.0, np.pi).astype(np.float32)


def directional_continuity(
    skeleton: np.ndarray,
    tangent: np.ndarray,
    lookahead: int,
    tolerance: int,
) -> Tuple[float, float]:
    coords = np.argwhere(skeleton)
    if coords.size == 0:
        return float("nan"), float("nan")
    lookahead = max(2, int(lookahead))
    tolerance = max(0, int(tolerance))
    support = ndimage.maximum_filter(skeleton.astype(np.uint8), size=2 * tolerance + 1, mode="constant") > 0
    y = coords[:, 0]
    x = coords[:, 1]
    theta = tangent[y, x]
    dy = np.sin(theta)
    dx = np.cos(theta)
    forward = np.zeros(len(coords), dtype=bool)
    backward = np.zeros(len(coords), dtype=bool)
    height, width = skeleton.shape

    for distance in range(2, lookahead + 1):
        fy = np.rint(y + distance * dy).astype(np.int64)
        fx = np.rint(x + distance * dx).astype(np.int64)
        by = np.rint(y - distance * dy).astype(np.int64)
        bx = np.rint(x - distance * dx).astype(np.int64)

        valid_f = (fy >= 0) & (fy < height) & (fx >= 0) & (fx < width)
        valid_b = (by >= 0) & (by < height) & (bx >= 0) & (bx < width)
        idx_f = np.flatnonzero(valid_f & ~forward)
        idx_b = np.flatnonzero(valid_b & ~backward)
        if idx_f.size:
            forward[idx_f] = support[fy[idx_f], fx[idx_f]]
        if idx_b.size:
            backward[idx_b] = support[by[idx_b], bx[idx_b]]
        if np.all(forward & backward):
            break

    mean_side_continuity = float(np.mean((forward.astype(np.float32) + backward.astype(np.float32)) / 2.0))
    bidirectional_ratio = float(np.mean(forward & backward))
    return mean_side_continuity, bidirectional_ratio


def mean_local_orientation_entropy(
    edge: np.ndarray,
    tangent: np.ndarray,
    bins: int,
    window: int,
    min_edges: int,
) -> float:
    edge = np.asarray(edge, dtype=bool)
    if not np.any(edge):
        return float("nan")
    bins = max(2, int(bins))
    window = odd_at_least_three(window)
    bin_index = np.floor(np.mod(tangent, np.pi) / np.pi * bins).astype(np.int32)
    bin_index = np.clip(bin_index, 0, bins - 1)
    area = float(window * window)
    total = uniform_filter(edge.astype(np.float32), size=window, mode="reflect") * area
    entropy = np.zeros_like(total, dtype=np.float32)
    for b in range(bins):
        indicator = (edge & (bin_index == b)).astype(np.float32)
        count = uniform_filter(indicator, size=window, mode="reflect") * area
        probability = count / (total + EPS)
        entropy -= np.where(probability > 0, probability * np.log(probability + EPS), 0.0)
    entropy /= math.log(float(bins))
    valid = edge & (total >= int(min_edges))
    if not np.any(valid):
        return float("nan")
    return float(np.mean(entropy[valid]))


def safe_corrcoef(a: np.ndarray, b: np.ndarray, mask: Optional[np.ndarray] = None) -> float:
    if mask is not None:
        a = a[mask]
        b = b[mask]
    else:
        a = a.ravel()
        b = b.ravel()
    finite = np.isfinite(a) & np.isfinite(b)
    a = a[finite]
    b = b[finite]
    if a.size < 3 or np.std(a) <= EPS or np.std(b) <= EPS:
        return float("nan")
    return float(np.corrcoef(a, b)[0, 1])


def overlap_metrics(edge: np.ndarray, target: np.ndarray) -> Dict[str, float]:
    edge = np.asarray(edge, dtype=bool)
    target = np.asarray(target, dtype=bool)
    intersection = float(np.count_nonzero(edge & target))
    union = float(np.count_nonzero(edge | target))
    edge_count = float(np.count_nonzero(edge))
    target_count = float(np.count_nonzero(target))
    precision = intersection / edge_count if edge_count > 0 else float("nan")
    recall = intersection / target_count if target_count > 0 else float("nan")
    iou = intersection / union if union > 0 else float("nan")
    f1 = 2.0 * precision * recall / (precision + recall + EPS) if np.isfinite(precision) and np.isfinite(recall) else float("nan")
    return {"precision": precision, "recall": recall, "iou": iou, "f1": f1}


def coherence_correspondence_metrics(
    response: np.ndarray,
    edge: np.ndarray,
    coherence: np.ndarray,
    valid_mask: np.ndarray,
    cfg: MetricConfig,
) -> Dict[str, float]:
    valid_values = coherence[valid_mask & np.isfinite(coherence)]
    if valid_values.size == 0:
        return {
            "low_coherence_threshold": float("nan"),
            "edge_in_low_coherence_precision": float("nan"),
            "low_coherence_edge_recall": float("nan"),
            "edge_low_coherence_iou": float("nan"),
            "edge_low_coherence_f1": float("nan"),
            "response_low_vs_other_mean_ratio": float("nan"),
            "response_vs_one_minus_coherence_corr": float("nan"),
            "edge_to_low_coherence_mean_distance_px": float("nan"),
        }

    threshold = float(np.percentile(valid_values, cfg.low_coherence_percentile))
    low = valid_mask & (coherence <= threshold)
    if cfg.low_coherence_tolerance_px > 0:
        low_eval = morphology.dilation(low, morphology.disk(cfg.low_coherence_tolerance_px))
    else:
        low_eval = low
    edge_eval = edge & valid_mask
    ov = overlap_metrics(edge_eval, low_eval)

    mean_low = float(np.mean(response[low])) if np.any(low) else float("nan")
    other = valid_mask & ~low
    mean_other = float(np.mean(response[other])) if np.any(other) else float("nan")
    mean_ratio = mean_low / (mean_other + EPS) if np.isfinite(mean_low) and np.isfinite(mean_other) else float("nan")
    corr = safe_corrcoef(response, 1.0 - coherence, valid_mask)
    distance_map = ndimage.distance_transform_edt(~low)
    mean_distance = float(np.mean(distance_map[edge_eval])) if np.any(edge_eval) and np.any(low) else float("nan")

    return {
        "low_coherence_threshold": threshold,
        "edge_in_low_coherence_precision": ov["precision"],
        "low_coherence_edge_recall": ov["recall"],
        "edge_low_coherence_iou": ov["iou"],
        "edge_low_coherence_f1": ov["f1"],
        "response_low_vs_other_mean_ratio": mean_ratio,
        "response_vs_one_minus_coherence_corr": corr,
        "edge_to_low_coherence_mean_distance_px": mean_distance,
    }


def fault_correspondence_metrics(
    edge: np.ndarray,
    fault_mask: Optional[np.ndarray],
    tolerance_px: int,
) -> Dict[str, float]:
    empty = {
        "edge_near_fault_precision": float("nan"),
        "fault_coverage_recall": float("nan"),
        "fault_edge_f1": float("nan"),
        "edge_to_fault_mean_distance_px": float("nan"),
    }
    if fault_mask is None or not np.any(fault_mask):
        return empty
    edge = np.asarray(edge, dtype=bool)
    fault = np.asarray(fault_mask, dtype=bool)
    radius = max(0, int(tolerance_px))
    if radius > 0:
        fault_dilated = morphology.dilation(fault, morphology.disk(radius))
        edge_dilated = morphology.dilation(edge, morphology.disk(radius))
    else:
        fault_dilated = fault
        edge_dilated = edge
    edge_count = float(np.count_nonzero(edge))
    fault_count = float(np.count_nonzero(fault))
    precision = float(np.count_nonzero(edge & fault_dilated)) / edge_count if edge_count > 0 else float("nan")
    recall = float(np.count_nonzero(fault & edge_dilated)) / fault_count if fault_count > 0 else float("nan")
    f1 = 2.0 * precision * recall / (precision + recall + EPS) if np.isfinite(precision) and np.isfinite(recall) else float("nan")
    distance_map = ndimage.distance_transform_edt(~fault)
    mean_distance = float(np.mean(distance_map[edge])) if edge_count > 0 else float("nan")
    return {
        "edge_near_fault_precision": precision,
        "fault_coverage_recall": recall,
        "fault_edge_f1": f1,
        "edge_to_fault_mean_distance_px": mean_distance,
    }


def compute_edge_metrics(
    response: np.ndarray,
    binary: np.ndarray,
    pre: Mapping[str, np.ndarray],
    fault_mask: Optional[np.ndarray],
    cfg: MetricConfig,
) -> Tuple[Dict[str, float], Dict[str, np.ndarray]]:
    binary = np.asarray(binary, dtype=bool)
    skeleton = morphology.skeletonize(binary)
    valid_mask = make_valid_signal_mask(pre["agc"], cfg)
    tangent = local_tangent_orientation(pre["orientation"])
    labels, lengths, pixel_counts = component_length_data(skeleton)
    component_count = int(lengths.size)
    total_length = float(np.sum(lengths)) if component_count else 0.0
    skeleton_pixels = int(np.count_nonzero(skeleton))
    degree = skeleton_degree(skeleton)
    endpoint_count = int(np.count_nonzero(skeleton & (degree == 1)))
    isolated_pixel_count = int(np.count_nonzero(skeleton & (degree == 0)))
    branch_count, crossing_count, junction_total = count_junction_clusters(skeleton)

    if component_count:
        short_ids = np.flatnonzero(lengths < float(cfg.short_edge_length_px)) + 1
        isolated_ids = np.flatnonzero(lengths <= float(cfg.isolated_max_length_px)) + 1
        short_pixel_count = int(np.count_nonzero(np.isin(labels, short_ids))) if short_ids.size else 0
        isolated_component_pixels = int(np.count_nonzero(np.isin(labels, isolated_ids))) if isolated_ids.size else 0
        short_component_count = int(short_ids.size)
        isolated_component_count = int(isolated_ids.size)
    else:
        short_pixel_count = 0
        isolated_component_pixels = 0
        short_component_count = 0
        isolated_component_count = 0

    continuity, bidirectional = directional_continuity(
        skeleton,
        tangent,
        cfg.continuity_lookahead_px,
        cfg.continuity_tolerance_px,
    )
    orientation_entropy = mean_local_orientation_entropy(
        binary,
        tangent,
        cfg.orientation_bins,
        cfg.orientation_entropy_window,
        cfg.orientation_entropy_min_edges,
    )

    edge_count = int(np.count_nonzero(binary))
    valid_area = int(np.count_nonzero(valid_mask))
    edge_valid_count = int(np.count_nonzero(binary & valid_mask))

    metrics: Dict[str, float] = {
        "edge_threshold_used": float("nan"),
        "edge_pixel_count": float(edge_count),
        "skeleton_pixel_count": float(skeleton_pixels),
        "component_count": float(component_count),
        "connected_component_mean_length_px": float(np.mean(lengths)) if component_count else 0.0,
        "connected_component_median_length_px": float(np.median(lengths)) if component_count else 0.0,
        "connected_component_p90_length_px": float(np.percentile(lengths, 90.0)) if component_count else 0.0,
        "total_skeleton_graph_length_px": total_length,
        "endpoint_count": float(endpoint_count),
        "break_rate_endpoint_per_px": endpoint_count / (total_length + EPS) if total_length > 0 else float("nan"),
        "breaks_per_100px": 100.0 * endpoint_count / (total_length + EPS) if total_length > 0 else float("nan"),
        "fragmentation_components_per_100px": 100.0 * component_count / (total_length + EPS) if total_length > 0 else float("nan"),
        "local_directional_continuity": continuity,
        "bidirectional_continuity_ratio": bidirectional,
        "edge_density_full": edge_count / float(binary.size),
        "edge_density_valid_signal": edge_valid_count / float(valid_area) if valid_area > 0 else float("nan"),
        "branch_point_count": float(branch_count),
        "crossing_structure_count": float(crossing_count),
        "junction_count_total": float(junction_total),
        "local_orientation_entropy": orientation_entropy,
        "spurious_short_edge_pixel_ratio": short_pixel_count / float(skeleton_pixels) if skeleton_pixels > 0 else float("nan"),
        "short_component_ratio": short_component_count / float(component_count) if component_count > 0 else float("nan"),
        "isolated_response_ratio": isolated_component_pixels / float(skeleton_pixels) if skeleton_pixels > 0 else float("nan"),
        "isolated_component_ratio": isolated_component_count / float(component_count) if component_count > 0 else float("nan"),
        "strict_isolated_pixel_ratio": isolated_pixel_count / float(skeleton_pixels) if skeleton_pixels > 0 else float("nan"),
        "mean_edge_response": float(np.mean(response[binary])) if edge_count > 0 else float("nan"),
    }
    metrics.update(
        coherence_correspondence_metrics(
            response=response,
            edge=binary,
            coherence=pre["coherence"],
            valid_mask=valid_mask,
            cfg=cfg,
        )
    )
    metrics.update(fault_correspondence_metrics(binary, fault_mask, cfg.fault_tolerance_px))

    maps = {
        "binary": binary,
        "skeleton": skeleton,
        "valid_mask": valid_mask,
        "tangent": tangent,
    }
    return metrics, maps


# -----------------------------------------------------------------------------
# Fault masks and outputs
# -----------------------------------------------------------------------------


def build_mask_index(mask_dir: Optional[Path], recursive: bool) -> Dict[str, Path]:
    if mask_dir is None:
        return {}
    if not mask_dir.exists():
        raise FileNotFoundError(f"Fault-mask directory not found: {mask_dir}")
    iterator = mask_dir.rglob("*") if recursive else mask_dir.glob("*")
    index: Dict[str, Path] = {}
    for path in iterator:
        if path.is_file() and path.suffix.lower() in IMAGE_EXTENSIONS:
            index.setdefault(path.stem, path)
    return index


def load_fault_mask(mask_path: Optional[Path], shape: Tuple[int, int]) -> Optional[np.ndarray]:
    if mask_path is None:
        return None
    with Image.open(mask_path) as im:
        arr = np.asarray(im.convert("L"))
    if arr.shape != shape:
        resampling = getattr(Image, "Resampling", Image)
        resized = Image.fromarray(arr).resize((shape[1], shape[0]), resample=resampling.NEAREST)
        arr = np.asarray(resized)
    return arr > 0


def metric_definition_rows() -> List[Dict[str, str]]:
    return [
        {"metric": "connected_component_mean_length_px", "definition": "骨架各8邻域连通分量的图长度均值；水平/垂直邻接计1，斜向邻接计√2。", "direction": "通常越大表示长连续边缘更多。"},
        {"metric": "break_rate_endpoint_per_px", "definition": "骨架端点数 / 骨架总图长度。这里将端点密度作为可重复的断裂率定义。", "direction": "越低通常越连续。"},
        {"metric": "local_directional_continuity", "definition": "沿原图局部切向，在前后两个方向的给定搜索距离内找到边缘支持的平均比例，范围0–1。", "direction": "越高越连续。"},
        {"metric": "bidirectional_continuity_ratio", "definition": "同时在局部切向前、后方向找到边缘支持的骨架像素比例。", "direction": "越高越连续。"},
        {"metric": "edge_density_full", "definition": "二值边缘像素数 / 全图像素数。", "direction": "无统一优劣，需结合噪声和覆盖度解释。"},
        {"metric": "edge_density_valid_signal", "definition": "有效纹理区内的二值边缘密度，降低空白区面积差异的影响。", "direction": "无统一优劣。"},
        {"metric": "branch_point_count", "definition": "骨架8邻域度数为3的相邻候选像素聚类后得到的Y型分支事件数。", "direction": "需结合地质结构解释。"},
        {"metric": "crossing_structure_count", "definition": "骨架分支候选聚类中最大邻域度数≥4的交叉事件数。", "direction": "过高可能表示噪声交叉，也可能是真实复杂结构。"},
        {"metric": "local_orientation_entropy", "definition": "在滑动窗口内统计边缘位置局部切向分布，计算归一化Shannon熵后取均值，范围0–1。", "direction": "高值表示方向更杂乱/多样；低值表示方向更集中。"},
        {"metric": "edge_in_low_coherence_precision", "definition": "边缘像素落入低coherence区（含容差膨胀）的比例。", "direction": "针对断层响应时通常越高对应性越强，但也可能偏向低相干噪声。"},
        {"metric": "low_coherence_edge_recall", "definition": "低coherence区被边缘覆盖的比例。", "direction": "越高表示覆盖更充分。"},
        {"metric": "response_vs_one_minus_coherence_corr", "definition": "有效区内edge response与(1−coherence)的Pearson相关系数。", "direction": "正值越大表示响应更偏向低相干区。"},
        {"metric": "edge_to_low_coherence_mean_distance_px", "definition": "边缘像素到最近低coherence像素的平均欧氏距离。", "direction": "越低表示空间对应更紧密。"},
        {"metric": "edge_near_fault_precision", "definition": "提供断层掩码时，边缘像素位于断层容差带内的比例。", "direction": "越高越好。"},
        {"metric": "fault_coverage_recall", "definition": "提供断层掩码时，断层像素在边缘容差带内被覆盖的比例。", "direction": "越高越好。"},
        {"metric": "spurious_short_edge_pixel_ratio", "definition": "长度小于short_edge_length_px的骨架分量所含像素 / 全部骨架像素。无人工标签时仅是“虚假短边缘”的形态代理。", "direction": "通常越低越好。"},
        {"metric": "isolated_response_ratio", "definition": "长度不超过isolated_max_length_px的孤立分量像素 / 全部骨架像素。", "direction": "通常越低越好。"},
    ]


def flatten_parameters(*configs: Any, extra: Optional[Mapping[str, Any]] = None) -> pd.DataFrame:
    rows: List[Dict[str, Any]] = []
    for cfg in configs:
        prefix = cfg.__class__.__name__
        for key, value in asdict(cfg).items():
            rows.append({"group": prefix, "parameter": key, "value": value})
    if extra:
        for key, value in extra.items():
            rows.append({"group": "runtime", "parameter": key, "value": value})
    return pd.DataFrame(rows)


def summarize_algorithms(df: pd.DataFrame) -> pd.DataFrame:
    identifier_columns = {
        "image",
        "image_path",
        "algorithm",
        "fault_mask_found",
        "height",
        "width",
    }
    numeric_columns = [
        c for c in df.select_dtypes(include=[np.number]).columns if c not in identifier_columns
    ]
    if not numeric_columns:
        return pd.DataFrame()
    grouped = df.groupby("algorithm", sort=False)[numeric_columns].agg(["mean", "std", "median"])
    grouped.columns = [f"{metric}_{stat}" for metric, stat in grouped.columns]
    grouped = grouped.reset_index()
    counts = df.groupby("algorithm", sort=False).size().rename("row_count").reset_index()
    return counts.merge(grouped, on="algorithm", how="left")


def style_excel(path: Path) -> None:
    try:
        from openpyxl import load_workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter

        wb = load_workbook(path)
        for ws in wb.worksheets:
            ws.freeze_panes = "A2"
            ws.auto_filter.ref = ws.dimensions
            for cell in ws[1]:
                cell.font = Font(bold=True)
                cell.fill = PatternFill("solid", fgColor="D9EAF7")
                cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            for col_idx, column_cells in enumerate(ws.columns, start=1):
                max_len = min(max(len(str(c.value)) if c.value is not None else 0 for c in column_cells), 45)
                ws.column_dimensions[get_column_letter(col_idx)].width = max(10, max_len + 2)
        wb.save(path)
    except Exception as exc:
        logging.warning("Excel styling skipped: %s", exc)


def write_outputs(
    records: List[Dict[str, Any]],
    output_dir: Path,
    pre_cfg: PreprocessConfig,
    det_cfg: DetectorConfig,
    metric_cfg: MetricConfig,
    runtime: Mapping[str, Any],
) -> Tuple[Path, Path, Path]:
    df = pd.DataFrame.from_records(records)
    summary = summarize_algorithms(df)
    definitions = pd.DataFrame(metric_definition_rows())
    parameters = flatten_parameters(pre_cfg, det_cfg, metric_cfg, extra=runtime)

    csv_path = output_dir / "seismic_edge_metrics.csv"
    summary_csv_path = output_dir / "algorithm_summary.csv"
    xlsx_path = output_dir / "seismic_edge_metrics.xlsx"
    df.to_csv(csv_path, index=False, encoding="utf-8-sig")
    summary.to_csv(summary_csv_path, index=False, encoding="utf-8-sig")
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        df.to_excel(writer, sheet_name="per_image", index=False)
        summary.to_excel(writer, sheet_name="algorithm_summary", index=False)
        definitions.to_excel(writer, sheet_name="metric_definitions", index=False)
        parameters.to_excel(writer, sheet_name="parameters", index=False)
    style_excel(xlsx_path)
    return xlsx_path, csv_path, summary_csv_path


# -----------------------------------------------------------------------------
# Main processing
# -----------------------------------------------------------------------------


def process_dataset(
    input_path: Path,
    output_dir: Path,
    fault_mask_dir: Optional[Path],
    recursive: bool,
    save_maps: bool,
    pre_cfg: PreprocessConfig,
    det_cfg: DetectorConfig,
    metric_cfg: MetricConfig,
) -> List[Dict[str, Any]]:
    image_paths = collect_image_paths(input_path, recursive)
    if not image_paths:
        raise RuntimeError(f"No supported images found under: {input_path}")
    mask_index = build_mask_index(fault_mask_dir, recursive)
    maps_root = ensure_dir(output_dir / "maps") if save_maps else None
    records: List[Dict[str, Any]] = []

    for image_path in tqdm(image_paths, desc="Images"):
        try:
            gray = load_grayscale_image(image_path, pre_cfg)
            pre = preprocess_image(gray, pre_cfg)
            mask_path = mask_index.get(image_path.stem)
            fault_mask = load_fault_mask(mask_path, gray.shape)
            detector_outputs = run_detectors(pre, det_cfg)

            if maps_root is not None:
                image_map_dir = ensure_dir(maps_root / image_path.stem)
                save_gray(image_map_dir / "00_gray.png", pre["gray"])
                save_gray(image_map_dir / "01_nlmeans.png", pre["nlmeans"])
                save_gray(image_map_dir / "02_sgs.png", pre["sgs"])
                save_gray(image_map_dir / "03_agc.png", pre["agc"])
                save_gray(image_map_dir / "04_coherence.png", pre["coherence"])
                if fault_mask is not None:
                    save_binary(image_map_dir / "05_fault_mask.png", fault_mask)

            for algorithm, result in detector_outputs.items():
                metrics, metric_maps = compute_edge_metrics(
                    response=result["response"],
                    binary=result["binary"],
                    pre=pre,
                    fault_mask=fault_mask,
                    cfg=metric_cfg,
                )
                metrics["edge_threshold_used"] = float(result["threshold"])
                record: Dict[str, Any] = {
                    "image": image_path.name,
                    "image_path": str(image_path.resolve()),
                    "algorithm": algorithm,
                    "height": int(gray.shape[0]),
                    "width": int(gray.shape[1]),
                    "fault_mask_found": bool(fault_mask is not None),
                    "fault_mask_path": str(mask_path.resolve()) if mask_path is not None else "",
                }
                record.update(metrics)
                records.append(record)

                if maps_root is not None:
                    image_map_dir = maps_root / image_path.stem
                    safe_algorithm = algorithm.lower()
                    save_gray(image_map_dir / f"{safe_algorithm}_response.png", result["response"])
                    save_binary(image_map_dir / f"{safe_algorithm}_binary.png", metric_maps["binary"])
                    save_binary(image_map_dir / f"{safe_algorithm}_skeleton.png", metric_maps["skeleton"])
                    save_edge_overlay(
                        image_map_dir / f"{safe_algorithm}_overlay.png",
                        pre["agc"],
                        metric_maps["binary"],
                    )

        except Exception as exc:
            logging.exception("Failed processing %s: %s", image_path, exc)
            records.append(
                {
                    "image": image_path.name,
                    "image_path": str(image_path.resolve()),
                    "algorithm": "ERROR",
                    "error": repr(exc),
                }
            )
    return records


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="Compare SLoG, LoG and Canny edge topology/continuity metrics on seismic profile images.",
        formatter_class=argparse.ArgumentDefaultsHelpFormatter,
    )
    parser.add_argument("--input", required=True, help="One image file or a folder containing seismic images.")
    parser.add_argument("--output", required=True, help="Output directory.")
    parser.add_argument("--fault-mask-dir", default=None, help="Optional fault-mask folder; masks are matched by filename stem.")
    parser.add_argument("--recursive", action="store_true", help="Search input and mask folders recursively.")
    parser.add_argument("--save-maps", action="store_true", help="Save preprocessing, response, binary, skeleton and overlay maps.")
    parser.add_argument("--log-level", default="INFO", choices=["DEBUG", "INFO", "WARNING", "ERROR"])

    parser.add_argument("--device", default="", help="auto, cpu, cuda, cuda:0, etc.")
    parser.add_argument("--amp", action="store_true", help="Use FP16 for SGS diffusion when CUDA is available.")
    parser.add_argument("--nl-h", "--nlmeans-h-scale", dest="nl_h", type=float, default=float(DEFAULT_SGS_PARAMS["nl_h"]))
    parser.add_argument("--sgs-sigma", type=float, default=float(DEFAULT_SGS_PARAMS["sigma"]))
    parser.add_argument("--sgs-anisotropy", type=float, default=float(DEFAULT_SGS_PARAMS["anisotropy"]))
    parser.add_argument("--sgs-iterations", type=int, default=int(DEFAULT_SGS_PARAMS["iterations"]))
    parser.add_argument("--sgs-base-kappa", type=float, default=float(DEFAULT_SGS_PARAMS["base_kappa"]))
    parser.add_argument("--sgs-gamma", type=float, default=float(DEFAULT_SGS_PARAMS["gamma"]))
    parser.add_argument("--sgs-option", type=int, default=int(DEFAULT_SGS_PARAMS["option"]))
    parser.add_argument("--lambda-perp-base", type=float, default=float(DEFAULT_SGS_PARAMS["lambda_perp_base"]))
    parser.add_argument("--lambda-par-base", type=float, default=float(DEFAULT_SGS_PARAMS["lambda_par_base"]))
    parser.add_argument("--coherence-influence", type=float, default=float(DEFAULT_SGS_PARAMS["coherence_influence"]))
    parser.add_argument("--agc-window", type=int, default=int(DEFAULT_AGC_PARAMS["window"]))
    parser.add_argument("--agc-eps", type=float, default=float(DEFAULT_AGC_PARAMS["eps"]))
    parser.add_argument("--agc-p1", "--agc-low-percentile", dest="agc_p1", type=float, default=float(DEFAULT_AGC_PARAMS["p1"]))

    parser.add_argument("--tex-layout", default="fault,agc,edge", help="Recorded as project metadata; detector input remains AGC.")
    parser.add_argument("--glog-sigma", "--slog-sigma", dest="glog_sigma", type=float, default=float(DEFAULT_DET_PARAMS["glog_sigma"]))
    parser.add_argument("--glog-anisotropy", "--slog-anisotropy", dest="glog_anisotropy", type=float, default=float(DEFAULT_DET_PARAMS["glog_anisotropy"]))
    parser.add_argument("--glog-coherence-influence", type=float, default=float(DEFAULT_DET_PARAMS["glog_coherence_influence"]))
    parser.add_argument("--glog-nangles", "--slog-n-angles", dest="glog_nangles", type=int, default=int(DEFAULT_DET_PARAMS["glog_nangles"]))
    parser.add_argument("--glog-size-factor", type=float, default=float(DEFAULT_DET_PARAMS["glog_size_factor"]))
    parser.add_argument("--glog-sharpness", "--slog-sharpness", dest="glog_sharpness", type=float, default=float(DEFAULT_DET_PARAMS["glog_sharpness"]))
    parser.add_argument("--glog-alpha", "--slog-alpha", dest="glog_alpha", type=float, default=float(DEFAULT_DET_PARAMS["glog_alpha"]))
    parser.add_argument("--glog-aniso-clip", "--slog-aniso-clip", dest="glog_aniso_clip", type=float, default=float(DEFAULT_DET_PARAMS["glog_aniso_clip"]))
    parser.add_argument("--glog-gamma-scale", "--slog-gamma-scale", dest="glog_gamma_scale", type=float, default=float(DEFAULT_DET_PARAMS["glog_gamma_scale"]))
    parser.add_argument("--glog-use-absolute-sigma", action="store_true", default=bool(DEFAULT_DET_PARAMS["glog_use_absolute_sigma"]))
    parser.add_argument("--glog-dist-mode", choices=["relative", "absolute"], default=str(DEFAULT_DET_PARAMS["glog_dist_mode"]))
    parser.add_argument("--edge-thr", "--slog-edge-threshold", dest="edge_thr", type=float, default=0.35)
    parser.add_argument(
        "--glog-absolute-response",
        action="store_true",
        help="Optional two-polarity absolute SLoG response. Default reproduces the source code's signed normalization.",
    )
    parser.add_argument("--log-sigma", type=float, default=float(DEFAULT_DET_PARAMS["log_sigma"]))
    parser.add_argument("--log-edge-threshold", type=float, default=-1.0, help="Negative value selects Otsu threshold.")
    parser.add_argument("--canny-sigma", type=float, default=float(DEFAULT_DET_PARAMS["canny_sigma"]))
    parser.add_argument("--canny-low", type=float, default=float(DEFAULT_DET_PARAMS["canny_low"]))
    parser.add_argument("--canny-high", type=float, default=float(DEFAULT_DET_PARAMS["canny_high"]))
    parser.add_argument(
        "--canny-float-input",
        action="store_true",
        help="Use normalized float AGC for Canny. Default reproduces the source code by converting AGC to uint8 first.",
    )
    parser.add_argument(
        "--response-hysteresis",
        action="store_true",
        help="Use low/high hysteresis when binarizing SLoG and LoG responses. Default uses the selected high threshold directly.",
    )
    parser.add_argument("--hysteresis-low-ratio", type=float, default=0.50)

    # metrics
    parser.add_argument("--short-edge-length", type=float, default=10.0)
    parser.add_argument("--isolated-max-length", type=float, default=2.0)
    parser.add_argument("--continuity-lookahead", type=int, default=7)
    parser.add_argument("--continuity-tolerance", type=int, default=1)
    parser.add_argument("--orientation-bins", type=int, default=18)
    parser.add_argument("--orientation-entropy-window", type=int, default=21)
    parser.add_argument("--low-coherence-percentile", type=float, default=20.0)
    parser.add_argument("--low-coherence-tolerance", type=int, default=2)
    parser.add_argument("--fault-tolerance", type=int, default=3)
    return parser


def build_effective_parameter_dicts(
    pre_cfg: PreprocessConfig,
    det_cfg: DetectorConfig,
) -> Tuple[Dict[str, Any], Dict[str, Any], Dict[str, Any]]:
    sgs_params = {
        "sigma": pre_cfg.sgs_sigma,
        "anisotropy": pre_cfg.sgs_anisotropy,
        "iterations": pre_cfg.sgs_iterations,
        "base_kappa": pre_cfg.sgs_base_kappa,
        "gamma": pre_cfg.sgs_gamma,
        "option": pre_cfg.sgs_option,
        "lambda_perp_base": pre_cfg.lambda_perp_base,
        "lambda_par_base": pre_cfg.lambda_par_base,
        "coherence_influence": pre_cfg.coherence_influence,
        "device": pre_cfg.device,
        "prefer_fp16": pre_cfg.prefer_fp16,
        "nl_h": pre_cfg.nl_h,
    }
    agc_params = {
        "window": pre_cfg.agc_window,
        "eps": pre_cfg.agc_eps,
        "p1": pre_cfg.agc_p1,
    }
    det_params = {
        "tex_layout": det_cfg.tex_layout,
        "glog_sigma": det_cfg.glog_sigma,
        "glog_anisotropy": det_cfg.glog_anisotropy,
        "glog_nangles": det_cfg.glog_nangles,
        "glog_sharpness": det_cfg.glog_sharpness,
        "glog_alpha": det_cfg.glog_alpha,
        "glog_aniso_clip": det_cfg.glog_aniso_clip,
        "glog_gamma_scale": det_cfg.glog_gamma_scale,
        "glog_dist_mode": det_cfg.glog_dist_mode,
        "glog_coherence_influence": det_cfg.glog_coherence_influence,
        "glog_size_factor": det_cfg.glog_size_factor,
        "glog_use_absolute_sigma": det_cfg.glog_use_absolute_sigma,
        "edge_thr": det_cfg.edge_thr,
        "log_sigma": det_cfg.log_sigma,
        "canny_sigma": det_cfg.canny_sigma,
        "canny_low": det_cfg.canny_low,
        "canny_high": det_cfg.canny_high,
    }
    return sgs_params, agc_params, det_params


def configs_from_args(args: argparse.Namespace) -> Tuple[PreprocessConfig, DetectorConfig, MetricConfig]:
    pre_cfg = PreprocessConfig(
        nl_h=args.nl_h,
        sgs_sigma=args.sgs_sigma,
        sgs_anisotropy=args.sgs_anisotropy,
        sgs_iterations=args.sgs_iterations,
        sgs_base_kappa=args.sgs_base_kappa,
        sgs_gamma=args.sgs_gamma,
        sgs_option=args.sgs_option,
        lambda_perp_base=args.lambda_perp_base,
        lambda_par_base=args.lambda_par_base,
        coherence_influence=args.coherence_influence,
        agc_window=args.agc_window,
        agc_eps=args.agc_eps,
        agc_p1=args.agc_p1,
        device=args.device,
        prefer_fp16=args.amp,
    )
    det_cfg = DetectorConfig(
        tex_layout=args.tex_layout,
        glog_sigma=args.glog_sigma,
        glog_anisotropy=args.glog_anisotropy,
        glog_coherence_influence=args.glog_coherence_influence,
        glog_nangles=args.glog_nangles,
        glog_size_factor=args.glog_size_factor,
        glog_sharpness=args.glog_sharpness,
        glog_alpha=args.glog_alpha,
        glog_aniso_clip=args.glog_aniso_clip,
        glog_gamma_scale=args.glog_gamma_scale,
        glog_use_absolute_sigma=args.glog_use_absolute_sigma,
        glog_dist_mode=args.glog_dist_mode,
        glog_absolute_response=args.glog_absolute_response,
        edge_thr=args.edge_thr,
        log_sigma=args.log_sigma,
        log_edge_threshold=args.log_edge_threshold,
        canny_sigma=args.canny_sigma,
        canny_low=args.canny_low,
        canny_high=args.canny_high,
        canny_use_uint8_input=not args.canny_float_input,
        use_response_hysteresis=args.response_hysteresis,
        hysteresis_low_ratio=args.hysteresis_low_ratio,
    )
    metric_cfg = MetricConfig(
        short_edge_length_px=args.short_edge_length,
        isolated_max_length_px=args.isolated_max_length,
        continuity_lookahead_px=args.continuity_lookahead,
        continuity_tolerance_px=args.continuity_tolerance,
        orientation_bins=args.orientation_bins,
        orientation_entropy_window=args.orientation_entropy_window,
        low_coherence_percentile=args.low_coherence_percentile,
        low_coherence_tolerance_px=args.low_coherence_tolerance,
        fault_tolerance_px=args.fault_tolerance,
    )
    return pre_cfg, det_cfg, metric_cfg


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    logging.basicConfig(
        level=getattr(logging, args.log_level),
        format="%(asctime)s | %(levelname)s | %(message)s",
    )
    input_path = Path(args.input).expanduser()
    output_dir = ensure_dir(Path(args.output).expanduser())
    fault_mask_dir = Path(args.fault_mask_dir).expanduser() if args.fault_mask_dir else None
    pre_cfg, det_cfg, metric_cfg = configs_from_args(args)
    sgs_params, agc_params, det_params = build_effective_parameter_dicts(pre_cfg, det_cfg)
    (output_dir / "effective_parameters.json").write_text(
        json.dumps(
            {
                "sgs_params": sgs_params,
                "agc_params": agc_params,
                "det_params": det_params,
            },
            ensure_ascii=False,
            indent=2,
            default=str,
        ),
        encoding="utf-8",
    )

    records = process_dataset(
        input_path=input_path,
        output_dir=output_dir,
        fault_mask_dir=fault_mask_dir,
        recursive=args.recursive,
        save_maps=args.save_maps,
        pre_cfg=pre_cfg,
        det_cfg=det_cfg,
        metric_cfg=metric_cfg,
    )
    runtime = {
        "input": str(input_path.resolve()),
        "output": str(output_dir.resolve()),
        "fault_mask_dir": str(fault_mask_dir.resolve()) if fault_mask_dir else "",
        "recursive": bool(args.recursive),
        "save_maps": bool(args.save_maps),
        "torch_available": bool(torch is not None),
        "torch_cuda_available": bool(torch is not None and torch.cuda.is_available()),
    }
    xlsx_path, csv_path, summary_path = write_outputs(
        records,
        output_dir,
        pre_cfg,
        det_cfg,
        metric_cfg,
        runtime,
    )
    print(f"Done. Rows: {len(records)}")
    print(f"Excel:   {xlsx_path}")
    print(f"CSV:     {csv_path}")
    print(f"Summary: {summary_path}")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
